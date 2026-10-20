#!/usr/bin/env python3
"""Package R-produced tables only; no biological/statistical calculations."""
import argparse, csv, hashlib, json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

def main():
    p=argparse.ArgumentParser();p.add_argument('report_directory',type=Path);a=p.parse_args()
    dest=a.report_directory.resolve()
    summary=json.loads((dest/'NUMERICAL_SUMMARY.json').read_text())
    assert summary['validated'] and summary['mapping_records']==434
    target=dest/'Table_S11_reference_mapping_QC.xlsx';assert not target.exists()
    wb=Workbook();wb.remove(wb.active)
    names=['sample_QC','class_overall_summary','paired_changes','paired_change_summary',
           'variant_counts','PCA_variance','new_genotype_missing','metric_definitions',
           'reporting_notes','provenance']
    expected={}
    numeric_headers={'clean_pairs','primary_records','mapped_primary','mapped_fraction',
        'properly_paired_primary_fraction','mapq20_primary_fraction','covered_bases_q20Q20',
        'covered_bases_ge4','covered_base_mean_depth','covered_base_median_depth_lower',
        'assembly_bases','assembly_breadth_ge1','assembly_breadth_ge4','whole_assembly_mean_depth',
        'n_individuals','mean','sd','median','min','max','sites','component','variance_fraction','missing_fraction'}
    for name in names:
        source=dest/(name+'.csv');rows=list(csv.reader(source.open(newline='')))
        assert rows and len(rows)>1
        ws=wb.create_sheet(name);expected[name]=(len(rows),len(rows[0]))
        for i,row in enumerate(rows):
            values=[]
            for col,value in enumerate(row):
                header=rows[0][col]
                if i>0 and value and (header in numeric_headers or header.endswith('_mongolicus_minus_nanus')):
                    try: value=float(value) if any(c in value for c in '.eE') else int(value)
                    except ValueError: pass
                values.append(value)
            ws.append(values)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        rule=Side(style='thin',color='000000')
        for row in ws:
            for cell in row:
                cell.font=Font(name='Times New Roman',size=11,bold=cell.row==1)
                cell.alignment=Alignment(horizontal='left' if cell.column==1 else 'center',vertical='top',wrap_text=True)
                if cell.row==1: cell.border=Border(top=rule,bottom=rule)
                elif cell.row==ws.max_row: cell.border=Border(bottom=rule)
        for j in range(1,ws.max_column+1):
            width=min(48,max(12,max(len(str(ws.cell(i,j).value or '')) for i in range(1,min(ws.max_row,50)+1))+2))
            ws.column_dimensions[get_column_letter(j)].width=width
        ws.print_options.horizontalCentered=True;ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1'
    wb.save(target)
    reopened=load_workbook(target,data_only=True)
    assert reopened.sheetnames==names
    for name,shape in expected.items():assert (reopened[name].max_row,reopened[name].max_column)==shape
    record={'validated':True,'file':str(target),'sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
            'sheets':expected,'figure_QA':'NOT_ASSESSED','manuscript_changes':'NOT_APPLIED'}
    (dest/'WORKBOOK_VALIDATION.json').write_text(json.dumps(record,indent=2)+'\n')
if __name__=='__main__': main()
