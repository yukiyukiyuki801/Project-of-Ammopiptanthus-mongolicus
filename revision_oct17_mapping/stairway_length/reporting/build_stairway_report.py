#!/usr/bin/env python3
"""Postmeasurement S7 extension; never runs BAM extraction or demographic inference."""
import argparse,copy,csv,hashlib,json,os,tempfile
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment,Border,Font,Side
HERE=Path(__file__).resolve().parent;BASE=HERE.parent
GROUPS={'all':'All individuals','population_core':'Core populations','population_peripheral':'Peripheral populations','adaptive_peripheral':'Adaptive individuals in peripheral populations','genetic_nonadaptive':'Nonadaptive genetic groups pooled','genetic_core':'Core genetic group','genetic_peripheral':'Peripheral genetic group','genetic_adaptive_1':'Adaptive genetic group 1','genetic_adaptive_2':'Adaptive genetic group 2'}
COLUMNS=['Group','Individuals','Input panel SNPs','Projected chromosomes','SNPs contributing to projected SFS','Archived input L (bp)','New mapped ACGT extent (bp)','Uncovered target SNPs']
CAPTION=('C. Archived demographic inputs and independently reconstructed mapped-region extents. '
 'Input panel size is the archived easySFS dictionary size; contributing SNPs have a positive probability of being polymorphic at the historical chromosome projection. '
 'The mapped extent counts unique ACGT reference bases in observed covered regions containing these SNPs. '
 'It is a new reconstruction, not automatically the historical input L or a validated callable denominator. '
 'No Stairway Plot rerun or recalibration was performed. '
 'Matching archived input files does not prove their correspondence to the original production run or plotted histories. '
 'The earlier RAD-locus ×150 description should therefore be read in light of these recovered numeric inputs and distinct mapped measurements.')
def sha(p):
 h=hashlib.sha256()
 with open(p,'rb') as f:
  for b in iter(lambda:f.read(8*1024*1024),b''):h.update(b)
 return h.hexdigest()
def result_map(results,registry):
 targets={t['target_bed']:t for t in registry['targets']};assert len(targets)==len(registry['targets'])
 observed={}
 for r in results:
  key=r['target_bed'];assert key not in observed,'Duplicate measurement target';assert key in targets,'Unexpected measured target';t=targets[key]
  assert r['membership_group']==t['membership_group'] and r['samples']==len(t['samples'])
  assert r['projection_chromosomes']==t['projection_chromosomes']
  assert r['target_snps']==t['positive_SFS_contributor_sites']
  assert r['covered_target_snps']+r['uncovered_target_snps']==r['target_snps']
  assert r['observed_ACGT_extent_bp']+r['ambiguous_reference_bp_excluded']==r['observed_reference_extent_bp']
  assert r['historical_L_replaced'] is False and r['callable_denominator_verified'] is False
  assert all(isinstance(r[k],int) and r[k]>=0 for k in ('samples','projection_chromosomes','target_snps','covered_target_snps','uncovered_target_snps','observed_ACGT_extent_bp','observed_reference_extent_bp','ambiguous_reference_bp_excluded'))
  observed[key]=r
 assert set(observed)==set(targets),'Missing measured target sets'
 return observed,targets

def build_rows(audit,registry,results):
 measured,targets=result_map(results,registry);collapsed={};provenance=[];used=set()
 for b in audit['blueprints']:
  candidates={s['target_set']['target_bed'] for s in b['matching_saved_sfs'] if 'target_set' in s};assert candidates,'Blueprint lacks recovered targets'
  normalized=[]
  for key in sorted(candidates):
   assert key in measured;used.add(key);t=targets[key];r=measured[key]
   assert t['projection_chromosomes']==b['projection_chromosomes']
   assert t['sha256']==next(s['target_set']['sha256'] for s in b['matching_saved_sfs'] if s.get('target_set',{}).get('target_bed')==key)
   # Equivalent separately named dictionaries can collapse only with identical target bytes,
   # sample membership, projection, input size, numerical L and measured outcomes.
   identity=(t['membership_group'],tuple(t['samples']),t['all_datadict_sites'],t['projection_chromosomes'],t['target_bed_sha256'],b['historical_L_bp'])
   row={COLUMNS[0]:GROUPS[t['membership_group']],COLUMNS[1]:len(t['samples']),COLUMNS[2]:t['all_datadict_sites'],COLUMNS[3]:t['projection_chromosomes'],COLUMNS[4]:t['positive_SFS_contributor_sites'],COLUMNS[5]:b['historical_L_bp'],COLUMNS[6]:r['observed_ACGT_extent_bp'],COLUMNS[7]:r['uncovered_target_snps']}
   normalized.append((identity,row))
  assert len({x[0] for x in normalized})==1,'Ambiguous blueprint-to-target correspondence'
  identity,row=normalized[0];assert all(v==row for _,v in normalized),'Equivalent targets have different measured results'
  if identity in collapsed:assert collapsed[identity]==row,'Conflicting duplicate blueprint measurement'
  else:collapsed[identity]=row
  provenance.append({'blueprint':b['blueprint'],'blueprint_sha256':b['blueprint_sha256'],'population_label_original':b['population'],'projection_chromosomes':b['projection_chromosomes'],'historical_L_bp':b['historical_L_bp'],'projected_segregating_sum':b['projected_segregating_sum'],'L_equals_round_SFS_times150':b['L_equals_round_SFS_times150'],'matching_target_beds':sorted(candidates),'matching_saved_sfs':b['matching_saved_sfs'],'reader_row':row,'historical_execution_or_plot_match_proven':False})
 assert used==set(targets),'Some measured target sets lack historical blueprint correspondence'
 order={label:i for i,label in enumerate(GROUPS.values())}
 rows=sorted(collapsed.values(),key=lambda x:(order[x[COLUMNS[0]]],x[COLUMNS[2]],x[COLUMNS[3]],x[COLUMNS[5]]))
 return rows,provenance

def cell_state(c):
 return (c.value,c.data_type,copy.copy(c.font),copy.copy(c.fill),copy.copy(c.border),copy.copy(c.alignment),c.number_format,copy.copy(c.protection),copy.copy(c.hyperlink),copy.copy(c.comment))
def check_preserved(before,after):
 assert before.sheetnames==after.sheetnames,'Worksheet set/order changed'
 count=0
 for bs in before:
  fs=after[bs.title];assert list(bs.merged_cells.ranges)==[x for x in fs.merged_cells.ranges if x.min_row<=bs.max_row]
  assert {k:dict(v) for k,v in bs.column_dimensions.items()}=={k:dict(v) for k,v in fs.column_dimensions.items()},'Existing column dimensions changed'
  for row in bs:
   for c in row:assert cell_state(c)==cell_state(fs.cell(c.row,c.column)),('Existing cell changed',bs.title,c.coordinate);count+=1
  for k,v in bs.row_dimensions.items():assert dict(v)==dict(fs.row_dimensions[k]),('Existing row dimensions changed',bs.title,k)
  if bs.title!='S7':assert (bs.max_row,bs.max_column)==(fs.max_row,fs.max_column)
 return count

def append_workbook(source,output,rows):
 from zipfile import ZipFile
 from lxml import etree as E
 import posixpath
 assert Path(source).resolve()!=Path(output).resolve(),'Original workbook must remain immutable'
 assert not Path(output).exists(),'Refusing overwrite'
 before=openpyxl.load_workbook(source);start=before['S7'].max_row+2;header=start+1;last=header+len(rows);note=last+2
 NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships';q=lambda n:'{'+NS+'}'+n
 with ZipFile(source) as z:
  work=E.fromstring(z.read('xl/workbook.xml'));rels=E.fromstring(z.read('xl/_rels/workbook.xml.rels'))
  rid=next(x.get('{'+RNS+'}id') for x in work.find(q('sheets')) if x.get('name')=='S7')
  target=next(x.get('Target') for x in rels if x.get('Id')==rid);part=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
  original_sheet=z.read(part);sheet=E.fromstring(original_sheet);styles=E.fromstring(z.read('xl/styles.xml'))
  fonts=styles.find(q('fonts'));borders=styles.find(q('borders'));xfs=styles.find(q('cellXfs'))
  def font(bold=False):
   f=E.Element(q('font'));E.SubElement(f,q('name'),val='Arial');E.SubElement(f,q('sz'),val='9')
   if bold:E.SubElement(f,q('b'))
   i=len(fonts);fonts.append(f);fonts.set('count',str(len(fonts)));return i
  regular=font();bold=font(True)
  def border(top=False,bottom=False):
   b=E.Element(q('border'))
   for name,on in [('left',False),('right',False),('top',top),('bottom',bottom),('diagonal',False)]:
    side=E.SubElement(b,q(name))
    if on:side.set('style','thin');E.SubElement(side,q('color'),rgb='FF000000')
   i=len(borders);borders.append(b);borders.set('count',str(len(borders)));return i
  empty=border();head=border(True,True);bottom=border(False,True)
  def style(fid,bid,align='left',numeric=False):
   x=E.Element(q('xf'),numFmtId='1' if numeric else '0',fontId=str(fid),fillId='0',borderId=str(bid),xfId='0',applyFont='1',applyBorder='1',applyAlignment='1');E.SubElement(x,q('alignment'),horizontal=align,vertical='center',wrapText='1');i=len(xfs);xfs.append(x);xfs.set('count',str(len(xfs)));return i
  section=style(bold,empty);hl=style(bold,head);hc=style(bold,head,'center');bl=style(regular,empty);bc=style(regular,empty,'center',True);fl=style(regular,bottom);fc=style(regular,bottom,'center',True);ns=style(regular,empty)
  data=sheet.find(q('sheetData'))
  def addrow(index,values,style_ids,height):
   r=E.SubElement(data,q('row'),r=str(index),ht=str(height),customHeight='1')
   for j,(value,sid) in enumerate(zip(values,style_ids),1):
    c=E.SubElement(r,q('c'),r=f'{openpyxl.utils.get_column_letter(j)}{index}',s=str(sid))
    if isinstance(value,int):E.SubElement(c,q('v')).text=str(value)
    else:c.set('t','inlineStr');E.SubElement(E.SubElement(c,q('is')),q('t')).text=str(value)
  addrow(start,['C. Recovered sequence-length inputs and mapped-region check'],[section],22)
  addrow(header,COLUMNS,[hl]+[hc]*7,46)
  for i,row in enumerate(rows,header+1):addrow(i,[row[k] for k in COLUMNS],[fl if i==last else bl]+[fc if i==last else bc]*7,30)
  addrow(note,[CAPTION],[ns],110)
  merges=sheet.find(q('mergeCells'))
  if merges is None:
   merges=E.Element(q('mergeCells'));after_data=list(sheet).index(data)+1
   # mergeCells follows sheetData and optional sheetProtection/protectedRanges/scenarios/autoFilter/sortState/dataConsolidate/customSheetViews.
   earlier={'sheetProtection','protectedRanges','scenarios','autoFilter','sortState','dataConsolidate','customSheetViews'}
   while after_data<len(sheet) and E.QName(sheet[after_data]).localname in earlier:after_data+=1
   sheet.insert(after_data,merges)
  E.SubElement(merges,q('mergeCell'),ref=f'A{note}:H{note}');merges.set('count',str(len(merges)))
  dimension=sheet.find(q('dimension'))
  if dimension is not None:dimension.set('ref',f'A1:H{note}')
  replacement={part:E.tostring(sheet,encoding='utf-8'),'xl/styles.xml':E.tostring(styles,encoding='utf-8')}
  with ZipFile(output,'w') as out:
   for entry in z.infolist():out.writestr(entry,replacement.get(entry.filename,z.read(entry.filename)))
 with ZipFile(source) as a,ZipFile(output) as b:
  untouched=[n for n in a.namelist() if n not in replacement]
  assert all(a.read(n)==b.read(n) for n in untouched),'Other workbook ZIP parts changed'
  # Preserve exact existing cell XML text, attributes and style references.
  oldroot=E.fromstring(original_sheet);newroot=E.fromstring(b.read(part))
  oldrows=oldroot.find(q('sheetData'));newrows=newroot.find(q('sheetData'))
  assert all(E.tostring(x)==E.tostring(y) for x,y in zip(oldrows,list(newrows)[:len(oldrows)]))
 reopened=openpyxl.load_workbook(output);checks=check_preserved(before,reopened)
 for i in range(header,last+1):
  for j in range(1,9):
   c=reopened['S7'].cell(i,j);assert c.border.left.style is None and c.border.right.style is None
   assert c.alignment.horizontal==('left' if j==1 else 'center')
   assert c.border.top.style==('thin' if i==header else None)
   assert c.border.bottom.style==('thin' if i in (header,last) else None)
   if i>header:assert c.value==rows[i-header-1][COLUMNS[j-1]]
 return {'original_cells_verified':checks,'new_section_first_row':start,'new_data_rows':len(rows),'three_rule_format':'PASS','all_existing_cells_and_styles_preserved':'PASS','all_unmodified_workbook_parts_byte_identical':len(untouched),'existing_S7_rows_XML_identical':True,'other_worksheets_preserved':'PASS'}

def write_csv(p,rows,fields):
 with open(p,'w',newline='') as f:
  w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)

def main():
 p=argparse.ArgumentParser();p.add_argument('--measurement-root',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True);p.add_argument('--workbook',type=Path,required=True);a=p.parse_args()
 assert not a.output_dir.exists(),'Use a new output directory';assert a.workbook.is_file()
 marker=a.measurement_root/'aggregate/MEASUREMENT.json';m=json.loads(marker.read_text());assert m['validated'] is True
 assert m['fingerprint']['all_target_reference_bases_validated'] is True,'Target/reference validation missing'
 for file,digest in m['outputs'].items():assert Path(file).is_file() and sha(file)==digest,('Measurement product changed',file)
 registry_path=BASE/'MEASUREMENT_REGISTRY.json';audit_path=BASE/'SAVED_INPUT_AUDIT.json'
 assert sha(registry_path)==m['fingerprint']['registry_sha256'],'Measured registry differs'
 results_path=a.measurement_root/'aggregate/RESULTS.json';assert str(results_path) in m['outputs'],'Results missing from validated product manifest'
 registry=json.loads(registry_path.read_text());audit=json.loads(audit_path.read_text());results=json.loads(results_path.read_text());assert results==m['results']
 assert len(audit['blueprints'])==23 and len(results)==15
 root=BASE.parents[1]
 for blueprint in audit['blueprints']:
  assert sha(root/blueprint['blueprint'])==blueprint['blueprint_sha256'],'Archived blueprint changed'
 rows,provenance=build_rows(audit,registry,results);assert len(provenance)==23
 a.output_dir.parent.mkdir(parents=True,exist_ok=True);tmp=Path(tempfile.mkdtemp(prefix=a.output_dir.name+'.partial.',dir=a.output_dir.parent))
 try:
  write_csv(tmp/'Table_S7_extension.csv',rows,COLUMNS)
  (tmp/'Table_S7_extension.json').write_text(json.dumps({'caption':CAPTION,'rows':rows,'numbered_table':'S7','new_numbered_table_created':False,'historical_L_replaced':False},indent=2)+'\n')
  (tmp/'Blueprint_provenance_23_rows.json').write_text(json.dumps(provenance,indent=2)+'\n')
  flat=[{k:r[k] for k in ('blueprint','blueprint_sha256','population_label_original','projection_chromosomes','historical_L_bp','projected_segregating_sum','L_equals_round_SFS_times150','historical_execution_or_plot_match_proven')}|{'matching_target_beds':' | '.join(r['matching_target_beds'])} for r in provenance];write_csv(tmp/'Blueprint_provenance_23_rows.csv',flat,list(flat[0]))
  qa=append_workbook(a.workbook,tmp/'Supplementary_tables_S7_extended.xlsx',rows)
  qa.update(status='PASS',blueprints=23,measured_target_sets=15,collapsed_reader_rows=len(rows),all_matches_unambiguous=True,source_workbook=str(a.workbook.resolve()),source_workbook_sha256=sha(a.workbook),measurement_marker_sha256=sha(marker),registry_sha256=sha(registry_path),saved_input_audit_sha256=sha(audit_path),report_script_sha256=sha(__file__),outputs={f.name:sha(f) for f in tmp.iterdir() if f.is_file()})
  (tmp/'REPORT_QA.json').write_text(json.dumps(qa,indent=2)+'\n');tmp.replace(a.output_dir)
 except BaseException:
  # Keep partial artifacts for diagnosis; never overwrite the accepted workbook.
  raise
 print(json.dumps({'status':'PASS','reader_rows':len(rows),'output_dir':str(a.output_dir)}))
if __name__=='__main__':main()
